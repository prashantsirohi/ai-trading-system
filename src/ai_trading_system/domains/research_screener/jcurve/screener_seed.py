from __future__ import annotations

import csv
import hashlib
import io
import json
import shutil
import sqlite3
import tempfile
from dataclasses import dataclass
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any, Iterable

import duckdb
import pandas as pd

from ai_trading_system.domains.fundamentals.screener_client import ScreenerClient
from ai_trading_system.domains.research_screener.store import content_hash

from .store import JCurveStore


@dataclass(frozen=True)
class ScreenerSeedPolicy:
    raw: dict[str, Any]
    version: str
    policy_hash: str

    @classmethod
    def load(cls, path: Path) -> "ScreenerSeedPolicy":
        raw = json.loads(Path(path).read_text(encoding="utf-8"))
        required = {
            "policy_version", "screen", "history", "thresholds",
            "accepted_dispositions", "supplemental_symbols", "authority",
        }
        if set(raw) != required:
            raise ValueError("Screener seed policy has an invalid top-level shape")
        if raw["history"].get("cross_basis_splicing") != "PROHIBITED":
            raise ValueError("Screener seed policy must prohibit cross-basis splicing")
        if raw["history"].get("missing_cwip") != "NOT_DISCLOSED":
            raise ValueError("Screener seed policy must preserve missing CWIP")
        return cls(raw=raw, version=str(raw["policy_version"]), policy_hash=content_hash(raw))


@dataclass(frozen=True)
class ScreenInput:
    path: Path
    screen_url: str
    query_text: str
    symbols: tuple[str, ...]
    retrieved_at: datetime
    acquisition_mode: str
    temporary: bool


class ScreenerHistoryReader:
    """Read-only accounting history used only for J-curve discovery."""

    METRICS = {
        "net_block", "capital_work_in_progress", "depreciation",
        "sales", "operating_profit",
    }

    def __init__(self, db_path: Path, policy: ScreenerSeedPolicy, *, as_of_date: date):
        self.db_path = Path(db_path)
        self.policy = policy
        self.as_of_date = as_of_date
        if not self.db_path.is_file():
            raise FileNotFoundError(f"Screener fundamentals DB not found: {self.db_path}")

    def symbols(self) -> tuple[str, ...]:
        with self._connect() as conn:
            rows = conn.execute(
                "SELECT DISTINCT upper(trim(symbol)) FROM screener_financials ORDER BY 1"
            ).fetchall()
        return tuple(str(row[0]) for row in rows if row[0])

    def classify(
        self, symbol: str, *, screen_member: bool, supplemental_member: bool = False,
    ) -> dict[str, Any]:
        symbol = symbol.upper().strip()
        with self._connect() as conn:
            basis, basis_reason, annual, quarters = self._select_basis(conn, symbol)
            market_cap = self._market_cap(conn, symbol)
        metrics = self._metrics(annual, quarters, market_cap)
        decision = self._decision(metrics)
        payload = {
            "symbol": symbol,
            "screen_member": bool(screen_member),
            "supplemental_member": bool(supplemental_member),
            "statement_basis": basis,
            "basis_resolution_reason": basis_reason,
            "annual_period_count": len(annual),
            "aligned_quarter_count": len(quarters),
            "metrics": metrics,
            **decision,
        }
        payload["decision_hash"] = content_hash(payload)
        return payload

    def _connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(f"file:{self.db_path}?mode=ro", uri=True)
        conn.row_factory = sqlite3.Row
        return conn

    def _select_basis(
        self, conn: sqlite3.Connection, symbol: str,
    ) -> tuple[str | None, str, list[dict[str, Any]], list[dict[str, Any]]]:
        required_annual = int(self.policy.raw["history"]["required_annual_periods"])
        required_quarters = int(self.policy.raw["history"]["required_quarterly_periods"])
        candidates: dict[str, tuple[list[dict[str, Any]], list[dict[str, Any]]]] = {}
        for basis in self.policy.raw["history"]["basis_order"]:
            candidates[str(basis)] = (
                self._annual(conn, symbol, str(basis)),
                self._quarters(conn, symbol, str(basis)),
            )
        consolidated = candidates.get("consolidated", ([], []))
        standalone = candidates.get("standalone", ([], []))
        consolidated_complete = (
            len(consolidated[0]) >= required_annual
            and len(consolidated[1]) >= required_quarters
        )
        standalone_complete = (
            len(standalone[0]) >= required_annual
            and len(standalone[1]) >= required_quarters
        )
        consolidated_latest = consolidated[1][0]["report_date"] if consolidated[1] else None
        standalone_latest = standalone[1][0]["report_date"] if standalone[1] else None
        if consolidated_complete and (
            standalone_latest is None or consolidated_latest >= standalone_latest
        ):
            return "consolidated", "consolidated_complete_history", *consolidated
        if standalone_complete:
            reason = (
                "standalone_fallback_insufficient_consolidated_history"
                if consolidated[0] or consolidated[1]
                else "standalone_only_complete_history"
            )
            return "standalone", reason, *standalone
        usable = [
            (basis, rows[0], rows[1]) for basis, rows in candidates.items()
            if rows[0] or rows[1]
        ]
        if not usable:
            return None, "no_financial_history", [], []
        basis, annual, quarters = max(
            usable,
            key=lambda item: (
                len(item[1]) >= required_annual,
                len(item[1]), len(item[2]), item[0] == "standalone",
            ),
        )
        return basis, "insufficient_history_best_single_basis", annual, quarters

    def _annual(
        self, conn: sqlite3.Connection, symbol: str, basis: str,
    ) -> list[dict[str, Any]]:
        rows = conn.execute(
            """
            SELECT report_date, metric_id, value, available_at
            FROM screener_financials
            WHERE symbol = ? AND statement_basis = ? AND period_type = 'annual'
              AND available_at <= ?
              AND metric_id IN ('net_block', 'capital_work_in_progress', 'depreciation')
            ORDER BY report_date DESC, metric_id
            """,
            [symbol, basis, self.as_of_date.isoformat()],
        ).fetchall()
        return self._pivot(rows)

    def _quarters(
        self, conn: sqlite3.Connection, symbol: str, basis: str,
    ) -> list[dict[str, Any]]:
        rows = conn.execute(
            """
            SELECT report_date, metric_id, value, available_at
            FROM screener_financials
            WHERE symbol = ? AND statement_basis = ? AND period_type = 'quarterly'
              AND available_at <= ?
              AND metric_id IN ('sales', 'operating_profit', 'depreciation')
            ORDER BY report_date DESC, metric_id
            """,
            [symbol, basis, self.as_of_date.isoformat()],
        ).fetchall()
        return [
            row for row in self._pivot(rows)
            if all(row.get(metric) is not None for metric in ("sales", "operating_profit", "depreciation"))
        ]

    @staticmethod
    def _pivot(rows: Iterable[sqlite3.Row]) -> list[dict[str, Any]]:
        by_date: dict[str, dict[str, Any]] = {}
        for row in rows:
            report_date = str(row["report_date"])
            item = by_date.setdefault(
                report_date,
                {"report_date": report_date, "available_at": str(row["available_at"])},
            )
            item[str(row["metric_id"])] = row["value"]
        return [by_date[key] for key in sorted(by_date, reverse=True)]

    def _market_cap(self, conn: sqlite3.Connection, symbol: str) -> float | None:
        row = conn.execute(
            """
            SELECT market_cap_cr FROM screener_company_snapshot
            WHERE symbol = ? AND as_of_date <= ?
            ORDER BY as_of_date DESC LIMIT 1
            """,
            [symbol, self.as_of_date.isoformat()],
        ).fetchone()
        return float(row[0]) if row and row[0] is not None else None

    @staticmethod
    def _growth(current: float | None, previous: float | None) -> float | None:
        if current is None or previous in (None, 0):
            return None
        return float(current) / float(previous) - 1.0

    @staticmethod
    def _capital_base(row: dict[str, Any]) -> float | None:
        net_block = row.get("net_block")
        cwip = row.get("capital_work_in_progress")
        if net_block is None or cwip is None:
            return None
        return float(net_block) + float(cwip)

    def _metrics(
        self,
        annual: list[dict[str, Any]],
        quarters: list[dict[str, Any]],
        market_cap: float | None,
    ) -> dict[str, Any]:
        current = annual[0] if annual else {}
        previous = annual[1] if len(annual) >= 2 else {}
        two_year = annual[2] if len(annual) >= 3 else {}
        four_year = annual[4] if len(annual) >= 5 else {}
        current_base = self._capital_base(current)
        previous_base = self._capital_base(previous)
        two_year_base = self._capital_base(two_year)
        four_year_base = self._capital_base(four_year)
        base_delta_2y = (
            current_base - two_year_base
            if current_base is not None and two_year_base is not None else None
        )
        metrics: dict[str, Any] = {
            "current_report_date": current.get("report_date"),
            "previous_report_date": previous.get("report_date"),
            "two_year_report_date": two_year.get("report_date"),
            "four_year_report_date": four_year.get("report_date"),
            "net_block_current": current.get("net_block"),
            "net_block_previous": previous.get("net_block"),
            "net_block_two_year": two_year.get("net_block"),
            "cwip_current": current.get("capital_work_in_progress"),
            "cwip_previous": previous.get("capital_work_in_progress"),
            "cwip_two_year": two_year.get("capital_work_in_progress"),
            "depreciation_current": current.get("depreciation"),
            "depreciation_previous": previous.get("depreciation"),
            "capital_base_current": current_base,
            "capital_base_previous": previous_base,
            "capital_base_two_year": two_year_base,
            "capital_base_four_year": four_year_base,
            "capital_base_2y_growth": self._growth(current_base, two_year_base),
            "capital_base_4y_growth": self._growth(current_base, four_year_base),
            "capital_base_delta_2y_cr": base_delta_2y,
            "capital_base_delta_2y_to_market_cap": self._growth(
                market_cap + base_delta_2y if market_cap is not None and base_delta_2y is not None else None,
                market_cap,
            ),
            "net_block_yoy_growth": self._growth(current.get("net_block"), previous.get("net_block")),
            "net_block_2y_growth": self._growth(current.get("net_block"), two_year.get("net_block")),
            "cwip_yoy_growth": self._growth(
                current.get("capital_work_in_progress"), previous.get("capital_work_in_progress")
            ),
            "cwip_share": (
                float(current["capital_work_in_progress"]) / current_base
                if current_base not in (None, 0) and current.get("capital_work_in_progress") is not None
                else None
            ),
            "annual_depreciation_growth": self._growth(
                current.get("depreciation"), previous.get("depreciation")
            ),
            "market_cap_cr": market_cap,
        }
        required_quarters = int(self.policy.raw["history"]["required_quarterly_periods"])
        if len(quarters) >= required_quarters:
            recent = quarters[:4]
            prior = quarters[4:8]
            totals = {
                metric: (
                    sum(float(row[metric]) for row in recent),
                    sum(float(row[metric]) for row in prior),
                )
                for metric in ("sales", "operating_profit", "depreciation")
            }
            current_margin = (
                totals["operating_profit"][0] / totals["sales"][0]
                if totals["sales"][0] else None
            )
            previous_margin = (
                totals["operating_profit"][1] / totals["sales"][1]
                if totals["sales"][1] else None
            )
            metrics.update({
                "latest_quarter": recent[0]["report_date"],
                "ttm_sales_growth": self._growth(*totals["sales"]),
                "ttm_operating_profit_growth": self._growth(*totals["operating_profit"]),
                "ttm_depreciation_growth": self._growth(*totals["depreciation"]),
                "ttm_margin_expansion_pp": (
                    (current_margin - previous_margin) * 100.0
                    if current_margin is not None and previous_margin is not None else None
                ),
            })
        else:
            metrics.update({
                "latest_quarter": quarters[0]["report_date"] if quarters else None,
                "ttm_sales_growth": None,
                "ttm_operating_profit_growth": None,
                "ttm_depreciation_growth": None,
                "ttm_margin_expansion_pp": None,
            })
        metrics["source_data_hash"] = content_hash(metrics)
        return metrics

    def _decision(self, metrics: dict[str, Any]) -> dict[str, Any]:
        threshold = self.policy.raw["thresholds"]
        reasons: list[str] = []
        lane_relative_2y = self._at_least(
            metrics["capital_base_2y_growth"], threshold["capital_base_2y_growth"]
        )
        lane_relative_4y = self._at_least(
            metrics["capital_base_4y_growth"], threshold["capital_base_4y_growth"]
        )
        lane_absolute = (
            self._at_least(
                metrics["capital_base_2y_growth"], threshold["absolute_lane_min_2y_growth"]
            )
            and self._at_least(
                metrics["capital_base_delta_2y_to_market_cap"],
                threshold["absolute_lane_delta_to_market_cap"],
            )
        )
        cwip_incomplete = (
            metrics["cwip_current"] is None or metrics["cwip_two_year"] is None
        )
        depreciation_growth = self._first(
            metrics["ttm_depreciation_growth"], metrics["annual_depreciation_growth"]
        )
        lane_missing_cwip = (
            cwip_incomplete
            and self._at_least(
                metrics["net_block_2y_growth"], threshold["missing_cwip_net_block_2y_growth"]
            )
            and self._at_least(
                depreciation_growth, threshold["missing_cwip_depreciation_growth"]
            )
        )
        accounting_visible = any((lane_relative_2y, lane_relative_4y, lane_absolute, lane_missing_cwip))
        for passed, code in (
            (lane_relative_2y, "CAPITAL_BASE_2Y_GROWTH"),
            (lane_relative_4y, "CAPITAL_BASE_4Y_GROWTH"),
            (lane_absolute, "ABSOLUTE_MATERIALITY"),
            (lane_missing_cwip, "MISSING_CWIP_NET_BLOCK_FALLBACK"),
        ):
            if passed:
                reasons.append(code)
        build = (
            accounting_visible
            and self._at_least(metrics["cwip_yoy_growth"], threshold["build_cwip_yoy_growth"])
            and self._at_least(metrics["cwip_share"], threshold["build_cwip_share"])
        )
        commissioning = (
            accounting_visible
            and self._at_least(
                metrics["net_block_yoy_growth"], threshold["commissioning_net_block_yoy_growth"]
            )
            and self._at_least(
                depreciation_growth, threshold["commissioning_depreciation_growth"]
            )
        )
        clean_transfer = (
            commissioning
            and metrics["cwip_yoy_growth"] is not None
            and metrics["cwip_yoy_growth"] <= float(threshold["clean_transfer_cwip_decline"])
        )
        ramp = (
            accounting_visible
            and self._at_least(metrics["ttm_sales_growth"], threshold["ramp_sales_growth"])
            and metrics["ttm_operating_profit_growth"] is not None
            and metrics["ttm_sales_growth"] is not None
            and metrics["ttm_operating_profit_growth"] > metrics["ttm_sales_growth"]
            and self._at_least(
                metrics["ttm_margin_expansion_pp"], threshold["ramp_margin_expansion_pp"]
            )
        )
        if build:
            reasons.append("CWIP_BUILD_SIGNAL")
        if commissioning:
            reasons.append("NET_BLOCK_DEPRECIATION_COMMISSIONING_SIGNAL")
        if clean_transfer:
            reasons.append("CWIP_TO_NET_BLOCK_TRANSFER")
        if ramp:
            reasons.append("OPERATING_LEVERAGE_ACTIVE")
        if ramp:
            disposition = "RAMP_ACTIVE"
        elif commissioning:
            disposition = "COMMISSIONING"
        elif build:
            disposition = "BUILD"
        elif accounting_visible:
            disposition = "CAPEX_VISIBLE_UNCLASSIFIED"
        else:
            disposition = "NOT_SELECTED"
            reasons.append("ACCOUNTING_SIGNAL_NOT_PROVEN")
        accepted = disposition in set(self.policy.raw["accepted_dispositions"])
        return {
            "accounting_visible": accounting_visible,
            "build_signal": build,
            "commissioning_signal": commissioning,
            "clean_transfer_signal": clean_transfer,
            "ramp_signal": ramp,
            "disposition": disposition,
            "accepted": accepted,
            "reason_codes": reasons,
        }

    @staticmethod
    def _at_least(value: float | None, threshold: float) -> bool:
        return value is not None and float(value) >= float(threshold)

    @staticmethod
    def _first(*values: float | None) -> float | None:
        return next((value for value in values if value is not None), None)


class ScreenerSeedService:
    def __init__(
        self,
        *,
        store_path: Path,
        output_root: Path,
        fundamentals_db: Path,
        policy: ScreenerSeedPolicy,
    ):
        self.store = JCurveStore(store_path)
        self.store_path = Path(store_path)
        self.output_root = Path(output_root)
        self.fundamentals_db = Path(fundamentals_db)
        self.policy = policy

    def run(
        self,
        *,
        as_of_date: date,
        screen_id: int,
        screen_export: Path | None = None,
        screener_client: ScreenerClient | None = None,
        supplemental_symbols: tuple[str, ...] = (),
    ) -> dict[str, Any]:
        configured_screen = int(self.policy.raw["screen"]["screen_id"])
        if int(screen_id) != configured_screen:
            raise ValueError(
                f"screen_id {screen_id} is not governed by policy {self.policy.version}"
            )
        started_at = datetime.now(UTC)
        screen_input = self._screen_input(
            as_of_date=as_of_date,
            screen_id=screen_id,
            screen_export=screen_export,
            screener_client=screener_client,
        )
        try:
            screen_rows = parse_screen_export(screen_input.path)
            raw_bytes = screen_input.path.read_bytes()
        finally:
            if screen_input.temporary:
                screen_input.path.unlink(missing_ok=True)
        # The frozen export is authoritative. Rendered-page links are retained in
        # ScreenInput only as acquisition diagnostics because pagination can expose
        # a different subset than the complete export.
        screen_symbols = {row["symbol"] for row in screen_rows if row.get("symbol")}
        if not screen_symbols:
            raise ValueError("Screener export contains no resolvable symbols")
        raw_hash = hashlib.sha256(raw_bytes).hexdigest()
        governed_supplemental = {
            str(symbol).upper().strip()
            for symbol in (*supplemental_symbols, *self.policy.raw["supplemental_symbols"])
            if str(symbol).strip()
        }
        history = ScreenerHistoryReader(
            self.fundamentals_db, self.policy, as_of_date=as_of_date,
        )
        available_symbols = set(history.symbols())
        universe = sorted(screen_symbols | governed_supplemental)
        missing_history_symbols = sorted(set(universe) - available_symbols)
        decisions = [
            (
                history.classify(
                    symbol,
                    screen_member=symbol in screen_symbols,
                    supplemental_member=symbol in governed_supplemental,
                )
                if symbol in available_symbols
                else self._history_unavailable_decision(
                    symbol,
                    screen_member=symbol in screen_symbols,
                    supplemental_member=symbol in governed_supplemental,
                )
            )
            for symbol in universe
        ]
        screen_metadata = {str(row["symbol"]): row for row in screen_rows}
        for row in decisions:
            metadata = screen_metadata.get(row["symbol"], {})
            row.update({
                "screen_exchange": metadata.get("listing_exchange"),
                "screen_listing_code": metadata.get("listing_code"),
                "screen_isin": metadata.get("isin"),
                "screen_nse_symbol": metadata.get("nse_symbol"),
                "screen_bse_code": metadata.get("bse_code"),
            })
        self._resolve_identities(decisions, as_of_date=as_of_date)
        name_by_symbol = {
            str(row["symbol"]): row.get("company_name")
            for row in screen_rows if row.get("symbol")
        }
        for row in decisions:
            row["company_name"] = name_by_symbol.get(row["symbol"]) or row.get("company_name")
            row["decision_hash"] = content_hash({
                key: value for key, value in row.items() if key != "decision_hash"
            })
        snapshot_hash = content_hash({
            "policy_hash": self.policy.policy_hash,
            "as_of_date": as_of_date,
            "screen_id": screen_id,
            "screen_content_hash": raw_hash,
            "query_hash": content_hash(screen_input.query_text),
            "decisions": [row["decision_hash"] for row in decisions],
        })
        run_id = f"jcurve-seed-{as_of_date}-{snapshot_hash[:16]}"
        output_dir = self.output_root / run_id
        if self.store.completed_seed_run(run_id) and output_dir.is_dir():
            return {
                "run_id": run_id, "status": "COMPLETED", "reused": True,
                "output_dir": str(output_dir),
            }
        if output_dir.exists():
            raise FileExistsError(f"incomplete J-curve seed output already exists: {output_dir}")
        output_dir.mkdir(parents=True)
        try:
            suffix = detect_screen_export_suffix(raw_bytes)
            source_path = output_dir / f"screener-screen-{screen_id}{suffix}"
            source_path.write_bytes(raw_bytes)
            artifact_id = f"artifact:jcurve-screener-screen:{raw_hash[:28]}"
            for row in decisions:
                row["source_artifact_id"] = artifact_id
            resolved_count = sum(row["identity_status"] == "RESOLVED" for row in decisions)
            candidate_count = sum(row["accepted"] and row["identity_status"] == "RESOLVED" for row in decisions)
            payload = {
                "run_id": run_id,
                "screen_id": screen_id,
                "screen_url": screen_input.screen_url,
                "as_of_date": as_of_date,
                "policy_version": self.policy.version,
                "policy_hash": self.policy.policy_hash,
                "query_text": screen_input.query_text,
                "query_hash": content_hash(screen_input.query_text),
                "source_artifact_id": artifact_id,
                "snapshot_hash": snapshot_hash,
                "source_row_count": len(screen_rows),
                "evaluated_count": len(decisions),
                "history_available_count": len(decisions) - len(missing_history_symbols),
                "history_unavailable_count": len(missing_history_symbols),
                "history_coverage_rate": (
                    (len(decisions) - len(missing_history_symbols)) / len(decisions)
                    if decisions else 0.0
                ),
                "resolved_count": resolved_count,
                "candidate_count": candidate_count,
                "started_at": started_at,
                "artifact": {
                    "artifact_id": artifact_id,
                    "ingestion_run_id": f"ingest:{run_id}:{artifact_id}",
                    "source_key": "screener_screen_export",
                    "provider": "SCREENER",
                    "source_url": screen_input.screen_url,
                    "local_dataset_id": f"screener.screen:{screen_id}:{as_of_date}",
                    "effective_date": as_of_date,
                    "published_at": None,
                    "retrieved_at": screen_input.retrieved_at,
                    "content_hash": raw_hash,
                    "byte_count": len(raw_bytes),
                    "row_count": len(screen_rows),
                    "parser_version": self.policy.version,
                    "schema_version": "jcurve-screener-seed-v1",
                    "validation_status": "VALID",
                    "metadata": {
                        "query_hash": content_hash(screen_input.query_text),
                        "acquisition_mode": screen_input.acquisition_mode,
                        "local_path": source_path.name,
                    },
                },
                "candidates": decisions,
            }
            self._write_outputs(output_dir, payload)
            self.store.persist_seed(payload)
            return {
                "run_id": run_id,
                "status": "COMPLETED",
                "reused": False,
                "output_dir": str(output_dir),
                "screen_row_count": len(screen_rows),
                "evaluated_count": len(decisions),
                "history_available_count": payload["history_available_count"],
                "history_unavailable_count": payload["history_unavailable_count"],
                "history_coverage_rate": payload["history_coverage_rate"],
                "resolved_count": resolved_count,
                "candidate_count": candidate_count,
                "unresolved_count": len(decisions) - resolved_count,
                "policy_version": self.policy.version,
            }
        except Exception:
            shutil.rmtree(output_dir)
            raise

    @staticmethod
    def _history_unavailable_decision(
        symbol: str, *, screen_member: bool, supplemental_member: bool,
    ) -> dict[str, Any]:
        return {
            "symbol": symbol,
            "screen_member": bool(screen_member),
            "supplemental_member": bool(supplemental_member),
            "statement_basis": None,
            "basis_resolution_reason": "screener_fundamentals_history_unavailable",
            "annual_period_count": 0,
            "aligned_quarter_count": 0,
            "metrics": {
                "history_available": False,
                "source_data_hash": content_hash({
                    "symbol": symbol,
                    "history_available": False,
                }),
            },
            "accounting_visible": False,
            "build_signal": False,
            "commissioning_signal": False,
            "clean_transfer_signal": False,
            "ramp_signal": False,
            "disposition": "HISTORY_UNAVAILABLE",
            "accepted": False,
            "reason_codes": ["SCREENER_FUNDAMENTALS_HISTORY_UNAVAILABLE"],
        }

    def _screen_input(
        self,
        *,
        as_of_date: date,
        screen_id: int,
        screen_export: Path | None,
        screener_client: ScreenerClient | None,
    ) -> ScreenInput:
        expected_query = str(self.policy.raw["screen"]["expected_query"])
        screen_url = str(self.policy.raw["screen"]["url"])
        if screen_export is not None:
            path = Path(screen_export)
            if not path.is_file() or path.stat().st_size == 0:
                raise FileNotFoundError(f"Screener screen export not found or empty: {path}")
            return ScreenInput(
                path=path,
                screen_url=screen_url,
                query_text=expected_query,
                symbols=(),
                retrieved_at=datetime.fromtimestamp(path.stat().st_mtime, tz=UTC),
                acquisition_mode="PROVIDED_EXPORT",
                temporary=False,
            )
        client = screener_client or ScreenerClient()
        with tempfile.TemporaryDirectory(prefix="jcurve-screener-") as temp_dir:
            target = Path(temp_dir) / f"screen-{screen_id}.download"
            result = client.download_screen_export(
                screen_id,
                destination=target,
                screen_url=screen_url,
                expected_query_terms=tuple(self.policy.raw["screen"]["required_query_terms"]),
            )
            if _canonical_query(result.query_text) != _canonical_query(expected_query):
                raise RuntimeError("Screener screen query drifted from the governed V1 query")
            with tempfile.NamedTemporaryFile(
                prefix=f"jcurve-screen-{screen_id}-", suffix=".download", delete=False,
            ) as handle:
                persisted = Path(handle.name)
            shutil.copyfile(result.path, persisted)
        return ScreenInput(
            path=persisted,
            screen_url=result.screen_url,
            query_text=result.query_text,
            symbols=result.symbols,
            retrieved_at=datetime.now(UTC),
            acquisition_mode="PLAYWRIGHT_AUTHENTICATED_EXPORT",
            temporary=True,
        )

    def _resolve_identities(self, rows: list[dict[str, Any]], *, as_of_date: date) -> None:
        conn = duckdb.connect(str(self.store_path), read_only=True)
        try:
            for row in rows:
                screen_isin = row.get("screen_isin")
                screen_exchange = row.get("screen_exchange")
                if screen_isin:
                    matches = conn.execute(
                        """
                        SELECT DISTINCT c.company_id, s.security_id, l.listing_id, c.legal_name
                        FROM listing_master l
                        JOIN security_master s ON s.security_id = l.security_id
                        JOIN company_master c ON c.company_id = s.company_id
                        WHERE s.isin = ? AND (? IS NULL OR l.exchange = ?)
                          AND l.valid_from <= ? AND (l.valid_to IS NULL OR l.valid_to >= ?)
                          AND s.valid_from <= ? AND (s.valid_to IS NULL OR s.valid_to >= ?)
                          AND c.valid_from <= ? AND (c.valid_to IS NULL OR c.valid_to >= ?)
                        ORDER BY c.company_id, s.security_id, l.listing_id
                        """,
                        [screen_isin, screen_exchange, screen_exchange, as_of_date, as_of_date,
                         as_of_date, as_of_date, as_of_date, as_of_date],
                    ).fetchall()
                else:
                    matches = conn.execute(
                        """
                    SELECT DISTINCT c.company_id, s.security_id, l.listing_id, c.legal_name
                    FROM listing_master l
                    JOIN security_master s ON s.security_id = l.security_id
                    JOIN company_master c ON c.company_id = s.company_id
                    WHERE l.exchange = 'NSE' AND upper(coalesce(l.symbol, '')) = ?
                      AND l.valid_from <= ? AND (l.valid_to IS NULL OR l.valid_to >= ?)
                      AND s.valid_from <= ? AND (s.valid_to IS NULL OR s.valid_to >= ?)
                      AND c.valid_from <= ? AND (c.valid_to IS NULL OR c.valid_to >= ?)
                    ORDER BY c.company_id, s.security_id, l.listing_id
                    """,
                    [row["symbol"], as_of_date, as_of_date, as_of_date, as_of_date,
                     as_of_date, as_of_date],
                    ).fetchall()
                if len(matches) == 1:
                    match = matches[0]
                    row.update({
                        "company_id": str(match[0]),
                        "security_id": str(match[1]),
                        "listing_id": str(match[2]),
                        "identity_status": "RESOLVED",
                        "company_name": str(match[3]),
                    })
                else:
                    row.update({
                        "company_id": None,
                        "security_id": None,
                        "listing_id": None,
                        "identity_status": "UNRESOLVED" if not matches else "AMBIGUOUS",
                    })
        finally:
            conn.close()

    @staticmethod
    def _write_outputs(output_dir: Path, payload: dict[str, Any]) -> None:
        result = {
            key: payload[key] for key in (
                "run_id", "screen_id", "screen_url", "as_of_date", "policy_version",
                "policy_hash", "query_hash", "snapshot_hash", "source_row_count",
                "evaluated_count", "history_available_count", "history_unavailable_count",
                "history_coverage_rate", "resolved_count", "candidate_count",
            )
        }
        (output_dir / "result.json").write_text(
            json.dumps(result, indent=2, default=str) + "\n", encoding="utf-8"
        )
        (output_dir / "manifest.json").write_text(
            json.dumps({
                **result,
                "query_text": payload["query_text"],
                "artifact": payload["artifact"],
                "decision_hashes": [row["decision_hash"] for row in payload["candidates"]],
            }, indent=2, default=str) + "\n",
            encoding="utf-8",
        )
        fields = [
            "symbol", "company_name", "screen_member", "supplemental_member",
            "screen_exchange", "screen_listing_code", "screen_isin",
            "screen_nse_symbol", "screen_bse_code",
            "identity_status", "company_id",
            "security_id", "statement_basis", "basis_resolution_reason", "accounting_visible",
            "build_signal", "commissioning_signal", "clean_transfer_signal", "ramp_signal",
            "disposition", "accepted", "reason_codes", "metrics", "decision_hash",
        ]
        with (output_dir / "candidates.csv").open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            for row in payload["candidates"]:
                writer.writerow({
                    key: json.dumps(row[key], sort_keys=True, default=str)
                    if key in {"reason_codes", "metrics"} else row.get(key)
                    for key in fields
                })


def profile_baseline(
    cohort: dict[str, Any], *, reader: ScreenerHistoryReader,
) -> dict[str, Any]:
    """Run the deterministic accounting model over an explicit calibration cohort."""

    available = set(reader.symbols())
    accepted_dispositions = set(reader.policy.raw["accepted_dispositions"])
    by_role: dict[str, dict[str, int]] = {}
    missing: list[str] = []
    accepted_count = 0
    for member in cohort["members"]:
        symbol = str(member["nse_symbol"]).upper().strip()
        if symbol not in available:
            missing.append(symbol)
            continue
        decision = reader.classify(
            symbol,
            screen_member=bool(member.get("matched_screen_ids")),
            supplemental_member=True,
        )
        role = str(member.get("case_role", "legacy_member"))
        counts = by_role.setdefault(role, {})
        disposition = str(decision["disposition"])
        counts[disposition] = counts.get(disposition, 0) + 1
        if disposition in accepted_dispositions:
            accepted_count += 1
    total = len(cohort["members"])
    return {
        "cohort_version": cohort["cohort_version"],
        "case_count": total,
        "history_available_count": total - len(missing),
        "history_unavailable_count": len(missing),
        "history_coverage_rate": (total - len(missing)) / total if total else 0.0,
        "accepted_count": accepted_count,
        "accepted_rate": accepted_count / total if total else 0.0,
        "missing_symbols": missing,
        "dispositions_by_role": by_role,
        "warning": "Cohort membership and model acceptance are not audited outcome labels.",
    }


def parse_screen_export(path: Path) -> list[dict[str, Any]]:
    raw = Path(path).read_bytes()
    if not raw:
        raise ValueError("Screener screen export is empty")
    if raw.startswith(b"PK\x03\x04"):
        frame = pd.read_excel(io.BytesIO(raw))
    else:
        try:
            frame = pd.read_csv(io.BytesIO(raw))
        except UnicodeDecodeError as exc:
            raise ValueError("Unsupported Screener screen export format") from exc
    if frame.empty:
        raise ValueError("Screener screen export contains no rows")
    normalized = {_normalize_column(column): column for column in frame.columns}
    nse_column = next(
        (normalized[key] for key in ("nsecode", "nsesymbol", "symbol", "ticker") if key in normalized),
        None,
    )
    bse_column = next(
        (normalized[key] for key in ("bsecode", "bsesymbol") if key in normalized),
        None,
    )
    isin_column = normalized.get("isincode") or normalized.get("isin")
    name_column = next(
        (normalized[key] for key in ("name", "company", "companyname") if key in normalized),
        None,
    )
    if nse_column is None and bse_column is None:
        hyperlink_symbols = _xlsx_hyperlink_symbols(raw) if raw.startswith(b"PK\x03\x04") else []
        if len(hyperlink_symbols) == len(frame):
            symbols: list[str | None] = hyperlink_symbols
        else:
            raise ValueError(
                "Screener screen export must contain an NSE/BSE code column or one company hyperlink per row"
            )
    else:
        symbols = []
        for _, item in frame.iterrows():
            nse_symbol = _clean_symbol(item[nse_column]) if nse_column is not None else None
            bse_code = _clean_bse_code(item[bse_column]) if bse_column is not None else None
            isin = _clean_symbol(item[isin_column]) if isin_column is not None else None
            symbols.append(nse_symbol or bse_code or (f"ISIN:{isin}" if isin else None))
    rows = []
    for index, symbol in enumerate(symbols):
        if not symbol:
            continue
        company_name = None
        if name_column is not None and pd.notna(frame.iloc[index][name_column]):
            company_name = str(frame.iloc[index][name_column]).strip()
        item = frame.iloc[index]
        nse_symbol = _clean_symbol(item[nse_column]) if nse_column is not None else None
        bse_code = _clean_bse_code(item[bse_column]) if bse_column is not None else None
        isin = _clean_symbol(item[isin_column]) if isin_column is not None else None
        listing_exchange = "NSE" if nse_symbol else "BSE" if bse_code else None
        listing_code = nse_symbol or bse_code
        rows.append({
            "symbol": symbol,
            "company_name": company_name,
            "listing_exchange": listing_exchange,
            "listing_code": listing_code,
            "isin": isin,
            "nse_symbol": nse_symbol,
            "bse_code": bse_code,
        })
    if not rows:
        raise ValueError("Screener screen export contains no valid listing identifiers")
    if len({row["symbol"] for row in rows}) != len(rows):
        raise ValueError("Screener screen export contains duplicate symbols")
    return rows


def _xlsx_hyperlink_symbols(raw: bytes) -> list[str]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []
    workbook = load_workbook(io.BytesIO(raw), read_only=False, data_only=True)
    sheet = workbook.active
    symbols: list[str] = []
    for row in sheet.iter_rows():
        found = None
        for cell in row:
            target = str(cell.hyperlink.target) if cell.hyperlink else ""
            if "/company/" not in target:
                continue
            parts = [part for part in target.split("/") if part]
            try:
                index = parts.index("company")
            except ValueError:
                continue
            if index + 1 < len(parts):
                found = _clean_symbol(parts[index + 1])
                break
        if found:
            symbols.append(found)
    return symbols


def detect_screen_export_suffix(raw: bytes) -> str:
    return ".xlsx" if raw.startswith(b"PK\x03\x04") else ".csv"


def _normalize_column(value: object) -> str:
    return "".join(character for character in str(value).casefold() if character.isalnum())


def _canonical_query(value: str) -> str:
    canonical = "".join(str(value).casefold().split())
    while _has_redundant_outer_parentheses(canonical):
        canonical = canonical[1:-1]
    return canonical


def _has_redundant_outer_parentheses(value: str) -> bool:
    if len(value) < 2 or value[0] != "(" or value[-1] != ")":
        return False
    depth = 0
    for index, character in enumerate(value):
        if character == "(":
            depth += 1
        elif character == ")":
            depth -= 1
            if depth < 0:
                return False
        if depth == 0 and index != len(value) - 1:
            return False
    return depth == 0


def _clean_symbol(value: object) -> str | None:
    if value is None or pd.isna(value):
        return None
    symbol = str(value).strip().upper()
    return symbol or None


def _clean_bse_code(value: object) -> str | None:
    if value is None or pd.isna(value):
        return None
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text if text.isdigit() else None


__all__ = [
    "ScreenInput",
    "ScreenerHistoryReader",
    "ScreenerSeedPolicy",
    "ScreenerSeedService",
    "detect_screen_export_suffix",
    "parse_screen_export",
    "profile_baseline",
]
